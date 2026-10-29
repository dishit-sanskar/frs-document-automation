import os
import zipfile
import json
import re
from docx import Document

def extract_flowchart_images(docx_path, output_dir="screenshots"):
    os.makedirs(output_dir, exist_ok=True)
    extracted = False
    
    with zipfile.ZipFile(docx_path, 'r') as z:
        image_files = [f for f in z.namelist() if f.startswith('word/media/')]
        if image_files:
            image_files.sort(key=lambda x: z.getinfo(x).file_size, reverse=True)
            main_img = image_files[0]
            
            out_path = os.path.join(output_dir, "to_be_flowchart.png")
            with open(out_path, "wb") as f:
                f.write(z.read(main_img))
            print(f"[Extractor] Saved flowchart diagram to: {out_path}")
            extracted = True
            
    if not extracted:
        print("[Extractor] Warning: No embedded media images found in the .docx file.")

def parse_docx_to_spec(docx_path, output_json="spec.json"):
    doc = Document(docx_path)
    
    spec_data = {
        "client_name": "Vogue International",
        "module_name": "Custom Module",
        "version_history": [
            {
                "version": "V 1.0",
                "issue_date": "21-10-2026",
                "description": "To Be Process Flow, Masters & Transactions",
                "created_by": "Dishit Dhanesha",
                "approved_by": "Sahil Patel"
            }
        ],
        "doctypes": [],
        "reports": []
    }

    current_doctype = None
    current_report = None
    in_chapter_2 = False
    in_chapter_3 = False

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if "Chapter 2" in text:
            in_chapter_2 = True
            in_chapter_3 = False
            continue
        elif "Chapter 3" in text:
            in_chapter_2 = False
            in_chapter_3 = True
            continue

        if in_chapter_2 and para.style.name.startswith("Heading"):
            clean_name = re.sub(r'^\d+(\.\d+)*\s*', '', text)
            clean_name = clean_name.replace("Master", "").strip()
            
            current_doctype = {
                "name": clean_name,
                "description": "",
                "mandatory_fields": [],
                "validation_rules": [],
                "fields": [],
                "dummy_records": []
            }
            spec_data["doctypes"].append(current_doctype)
        elif in_chapter_2 and current_doctype:
            if "mandatory" in text.lower():
                fields = [f.strip() for f in text.split(":")[-1].split(",") if f.strip()]
                current_doctype["mandatory_fields"].extend(fields)
            else:
                current_doctype["description"] += f" {text}"

        if in_chapter_3 and para.style.name.startswith("Heading"):
            clean_rep = re.sub(r'^\d+(\.\d+)*\s*', '', text).strip()
            current_report = {
                "name": clean_rep,
                "fields_mapping": []
            }
            spec_data["reports"].append(current_report)

    for table in doc.tables:
        if in_chapter_3 and current_report:
            for row in table.rows[1:]:
                cells = [cell.text.strip() for cell in row.cells]
                if len(cells) >= 3:
                    mapping = {
                        "field_name": cells[0],
                        "doctype_name": cells[1] if len(cells) > 1 else "",
                        "fieldname": cells[2] if len(cells) > 2 else "",
                        "remark": cells[3] if len(cells) > 3 else ""
                    }
                    current_report["fields_mapping"].append(mapping)

    with open(output_json, "w", encoding="utf-8") as f:
        json.dump(spec_data, f, indent=2)
    print(f"[Parser] Generated spec configuration: {output_json}")

    extract_flowchart_images(docx_path)

if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        parse_docx_to_spec(sys.argv[1])
    else:
        print("Usage: python parse_docx_to_spec.py <path_to_brd_file.docx>")
