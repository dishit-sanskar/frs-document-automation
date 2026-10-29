import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_sanskar_frs_docx(data, screenshots_dir, output_filename):
    doc = Document()
    
    COLOR_RED = RGBColor(192, 0, 0)
    COLOR_BLACK = RGBColor(0, 0, 0)
    COLOR_GRAY = RGBColor(128, 128, 128)

    # 1. Header / Title
    title_p = doc.add_paragraph()
    run_title = title_p.add_run("Functional Requirement Document\n")
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_RED

    run_sub = title_p.add_run(f"PREPARED FOR:\n{data.get('client_name', 'Vogue International')}\n\n")
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = COLOR_BLACK

    run_submitted = title_p.add_run("Submitted by:\nSanskar Technolab Pvt. Ltd\nB-305 to 308, Shilp Corporate Park, Rajpath Rangoli Road, Bodakdev, Ahmedabad.\nwww.sanskartechnolab.com\n")
    run_submitted.font.size = Pt(10)
    run_submitted.font.color.rgb = COLOR_GRAY

    doc.add_page_break()

    # 2. Document Version History
    doc.add_heading("Document Version History", level=1)
    
    table_history = doc.add_table(rows=1, cols=5)
    table_history.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table_history.rows[0].cells
    hdr_titles = ["Document Version", "Issue Date", "Description", "Created By", "Approved By"]
    for i, title in enumerate(hdr_titles):
        hdr_cells[i].text = title
        
    for ver in data.get("version_history", []):
        row_cells = table_history.add_row().cells
        row_cells[0].text = ver.get("version", "V 1.0")
        row_cells[1].text = ver.get("issue_date", "")
        row_cells[2].text = ver.get("description", "")
        row_cells[3].text = ver.get("created_by", "")
        row_cells[4].text = ver.get("approved_by", "")

    doc.add_page_break()

    # 3. Chapter 1 Flowcharts
    doc.add_heading("Chapter 1 Flowcharts", level=1)
    doc.add_heading("1.1 To Be Process Flow", level=2)
    
    flowchart_path = os.path.join(screenshots_dir, "to_be_flowchart.png")
    if os.path.exists(flowchart_path):
        doc.add_picture(flowchart_path, width=Inches(6.0))

    doc.add_page_break()

    # 4. Chapter 2 Masters & Transactions
    doc.add_heading("Chapter 2 Masters & Transactions", level=1)
    
    for idx, doctype in enumerate(data.get("doctypes", []), start=1):
        dt_name = doctype.get("name")
        doc.add_heading(f"2.{idx} {dt_name}", level=2)

        list_img = os.path.join(screenshots_dir, f"{dt_name.lower().replace(' ', '_')}_list_view.png")
        if os.path.exists(list_img):
            doc.add_picture(list_img, width=Inches(6.0))

        form_img = os.path.join(screenshots_dir, f"{dt_name.lower().replace(' ', '_')}_form_view.png")
        if os.path.exists(form_img):
            doc.add_picture(form_img, width=Inches(6.0))

        doc.add_paragraph(doctype.get("description", ""))

        if doctype.get("mandatory_fields"):
            doc.add_paragraph(f"Mandatory Fields: {', '.join(doctype.get('mandatory_fields'))}", style='List Bullet')
        if doctype.get("validation_rules"):
            p = doc.add_paragraph("Validation Rules:", style='List Bullet')
            for rule in doctype.get("validation_rules"):
                doc.add_paragraph(rule, style='List Bullet 2')

    doc.add_page_break()

    # 5. Chapter 3 Reports Mapping
    doc.add_heading("Chapter 3 Reports Mapping", level=1)
    
    for idx, report in enumerate(data.get("reports", []), start=1):
        rep_name = report.get("name")
        doc.add_heading(f"3.{idx} {rep_name}", level=2)

        rep_img = os.path.join(screenshots_dir, f"{rep_name.lower().replace(' ', '_')}_report.png")
        if os.path.exists(rep_img):
            doc.add_picture(rep_img, width=Inches(6.0))

        table_rep = doc.add_table(rows=1, cols=4)
        table_rep.alignment = WD_TABLE_ALIGNMENT.CENTER
        r_hdr = table_rep.rows[0].cells
        r_hdr[0].text = "Field Name"
        r_hdr[1].text = "ERPNext Doctype Name"
        r_hdr[2].text = "ERPNext Field"
        r_hdr[3].text = "Remark"

        for field in report.get("fields_mapping", []):
            row_c = table_rep.add_row().cells
            row_c[0].text = field.get("field_name", "")
            row_c[1].text = field.get("doctype_name", "")
            row_c[2].text = field.get("fieldname", "")
            row_c[3].text = field.get("remark", "")

    doc.save(output_filename)
    print(f"Document successfully created: {output_filename}")
